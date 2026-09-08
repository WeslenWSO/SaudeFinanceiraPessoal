"""Montagem de contexto para relatórios de impressão de lote."""

from __future__ import annotations

from collections import OrderedDict
from decimal import Decimal

from django.db.models import Max, Min
from django.utils import timezone

from empresa.models import Empresa

from .models import FaturamentoMedico, ItemServico, Lote

CONVENIOS_LAYOUT_PUBLICO_KEYWORDS = (
    'FUSEX',
    'POLICIA MILITAR',
    'POLÍCIA MILITAR',
    'CORPO DE BOMBEIRO',
    'BOMBEIRO',
    'PP SAUDE',
    'PP SAÚDE',
    'GEAP',
    'POSTAL',
)

CONVENIOS_RELATORIO_COLUNA_GUIA = (
    'FUSEX',
    'POLICIA MILITAR',
    'POLÍCIA MILITAR',
)

MESES_PT = (
    '',
    'JANEIRO',
    'FEVEREIRO',
    'MARÇO',
    'ABRIL',
    'MAIO',
    'JUNHO',
    'JULHO',
    'AGOSTO',
    'SETEMBRO',
    'OUTUBRO',
    'NOVEMBRO',
    'DEZEMBRO',
)

# Ordem do resumo no modelo PP Saúde / convênios públicos
RESUMO_MODALIDADES_PUBLICO = (
    ('MG', 'QUANTIDADE DE MAMOGRAFIA'),
    ('CT', 'QUANTIDADE DE TOMOGRAFIA'),
    ('US', 'QUANTIDADE DE ULTRASSONOGRAFIA'),
    ('CR', 'QUANTIDADE DE RAIO X'),
    ('MR', 'QUANTIDADE DE RESSONÂNCIA'),
    ('EG', 'QUANTIDADE DE ELETROENCEFALOGRAMA'),
    ('EC', 'QUANTIDADE DE ELETROCARDIOGRAMA'),
)


def convenio_usa_layout_publico(nome_convenio: str) -> bool:
    nome = (nome_convenio or '').upper()
    return any(palavra in nome for palavra in CONVENIOS_LAYOUT_PUBLICO_KEYWORDS)


def convenio_relatorio_coluna_guia(nome_convenio: str) -> bool:
    """FUSEX e PM exibem número da guia além do associado."""
    nome = (nome_convenio or '').upper().strip()
    if any(palavra in nome for palavra in CONVENIOS_RELATORIO_COLUNA_GUIA):
        return True
    if nome == 'PM' or nome.startswith('PM ') or nome.endswith(' PM') or ' PM ' in f' {nome} ':
        return True
    return False


def _validar_acesso_lote(lote_id, empresa_id):
    lote = Lote.objects.get(id=lote_id)
    if lote.empresa_id != int(empresa_id):
        raise PermissionError('Acesso negado')
    return lote


def _validar_acesso_lotes(lote_ids, empresa_id):
    ids = [int(i) for i in lote_ids]
    lotes = list(Lote.objects.filter(id__in=ids, empresa_id=int(empresa_id)).order_by('-id'))
    if len(lotes) != len(ids):
        raise Lote.DoesNotExist
    return lotes


def _modalidade_item(faturamento, item=None):
    if item and item.modalidade:
        return item.modalidade
    obs = faturamento.observacao or ''
    if 'Modalidade:' in obs:
        for parte in obs.splitlines():
            if parte.strip().lower().startswith('modalidade:'):
                valor = parte.split(':', 1)[-1].strip()
                if valor:
                    return valor
    return _inferir_modalidade(item.servico if item else faturamento.servico, '')


def _inferir_modalidade(procedimento, modalidade):
    mod = (modalidade or '').strip().upper()
    if mod and mod != '-':
        if mod == 'RX':
            return 'CR'
        return mod
    p = (procedimento or '').lower()
    if p.startswith('rm ') or 'resson' in p:
        return 'MR'
    if p.startswith('tc ') or 'tomograf' in p:
        return 'CT'
    if (
        p.startswith('us ')
        or 'ultrassom' in p
        or 'ultrassonograf' in p
        or 'doppler' in p
    ):
        return 'US'
    if p.startswith('rx ') or ' raio' in p or p.startswith('raio'):
        return 'CR'
    if 'mamograf' in p:
        return 'MG'
    if 'eletrocardiograma' in p or p.startswith('ecg'):
        return 'EC'
    if 'eletroencefalograma' in p or p.startswith('eeg'):
        return 'EG'
    return '-'


def _normalizar_codigo_modalidade(codigo):
    mod = (codigo or '').strip().upper()
    if mod == 'RX':
        return 'CR'
    return mod


def _montar_resumo_publico(linhas):
    contagem = {codigo: 0 for codigo, _ in RESUMO_MODALIDADES_PUBLICO}
    quantidade_total = 0
    valor_total = Decimal('0')
    for linha in linhas:
        quantidade_total += 1
        valor_total += linha.get('valor') or Decimal('0')
        codigo = _normalizar_codigo_modalidade(linha.get('modalidade'))
        if codigo in contagem:
            contagem[codigo] += 1
    resumo_modalidades = [
        {'codigo': codigo, 'label': label, 'quantidade': contagem[codigo]}
        for codigo, label in RESUMO_MODALIDADES_PUBLICO
    ]
    return {
        'modalidades': resumo_modalidades,
        'quantidade_total': quantidade_total,
        'valor_total': valor_total,
    }


def _mes_referencia_label(data_ref):
    if not data_ref:
        return ''
    mes = MESES_PT[data_ref.month] if 1 <= data_ref.month <= 12 else ''
    return f'{mes} {data_ref.year}' if mes else str(data_ref.year)


def _local_empresa(empresa):
    """Primeira linha do endereço da empresa para o campo Local do rodapé."""
    endereco = (empresa.endereco or '').strip()
    if not endereco:
        return ''
    return endereco.splitlines()[0].strip()


def _protocolo_lote(lote: Lote) -> str:
    extrato = lote.linhas_extrato_pagamento.first()
    if extrato and (extrato.protocolo or '').strip():
        return extrato.protocolo.strip()
    prot = (
        FaturamentoMedico.objects.filter(lote=str(lote.id))
        .exclude(guia_lancada__isnull=True)
        .exclude(guia_lancada='')
        .values_list('guia_lancada', flat=True)
        .first()
    )
    return (prot or '').strip()


def _lote_convenio_lote(lote: Lote) -> str:
    extrato = lote.linhas_extrato_pagamento.first()
    if extrato and (extrato.lote or '').strip():
        return extrato.lote.strip()
    return ''


def _montar_secao_lote(lote: Lote, *, layout: str):
    chave = str(lote.id)
    faturamentos = FaturamentoMedico.objects.filter(lote=chave).order_by('data', 'guia', 'nome')
    items = (
        ItemServico.objects.filter(faturamento__in=faturamentos)
        .select_related('faturamento')
        .order_by('faturamento__data', 'faturamento__nome', 'faturamento__guia', 'id')
    )
    periodo_inicio = faturamentos.aggregate(min_data=Min('data'))['min_data']
    periodo_fim = faturamentos.aggregate(max_data=Max('data'))['max_data']
    total_geral = Decimal('0')
    resumo_publico = None

    if layout == 'publico':
        linhas = []
        for item in items:
            fat = item.faturamento
            valor_item = item.total if item.total is not None else (item.valor or Decimal('0'))
            modalidade = _modalidade_item(fat, item)
            linhas.append({
                'data': fat.data,
                'paciente': fat.nome or '-',
                'nome_associado': fat.nome_associado or fat.nome or '-',
                'numero_guia': (fat.guia or '').strip() or '-',
                'procedimento': item.servico or '-',
                'modalidade': modalidade,
                'com_contraste': item.com_contraste,
                'valor': valor_item,
            })
            total_geral += valor_item
        resumo_publico = _montar_resumo_publico(linhas)
        grouped_rows = linhas
        grouped_items = OrderedDict()
    else:
        grouped_items = OrderedDict()
        for item in items:
            beneficiario = item.faturamento.nome or 'Sem Nome'
            grouped_items.setdefault(beneficiario, []).append(item)
            total_geral += item.total or Decimal('0')
        grouped_rows = []
        linhas = []

    return {
        'lote': lote,
        'protocolo': _protocolo_lote(lote),
        'lote_convenio': _lote_convenio_lote(lote),
        'periodo_inicio': periodo_inicio,
        'periodo_fim': periodo_fim,
        'mes_referencia': _mes_referencia_label(periodo_fim or lote.data_lote),
        'grouped_items': grouped_items,
        'linhas': grouped_rows if layout == 'publico' else linhas,
        'resumo_publico': resumo_publico,
        'total_geral': total_geral,
    }


def montar_contexto_relatorio_lote(lote_id, empresa_id, *, layout='padrao', lote_ids=None):
    from .lote_utils import parse_lote_ids

    ids = parse_lote_ids(lote_ids) if lote_ids else parse_lote_ids(lote_id)
    if not ids:
        raise Lote.DoesNotExist
    lotes = _validar_acesso_lotes(ids, empresa_id)
    lote = lotes[0]
    empresa = Empresa.objects.get(id=empresa_id)

    secoes = [_montar_secao_lote(lt, layout=layout) for lt in lotes]
    primeira = secoes[0]
    total_geral = sum((s['total_geral'] for s in secoes), Decimal('0'))
    periodo_inicio = min(
        (s['periodo_inicio'] for s in secoes if s['periodo_inicio']),
        default=None,
    )
    periodo_fim = max(
        (s['periodo_fim'] for s in secoes if s['periodo_fim']),
        default=None,
    )

    return {
        'lote': lote,
        'secoes': secoes,
        'empresa': empresa,
        'convenio_nome': lote.convenio or '',
        'local_relatorio': _local_empresa(empresa),
        'periodo_inicio': periodo_inicio,
        'periodo_fim': periodo_fim,
        'mes_referencia': _mes_referencia_label(periodo_fim or lote.data_lote),
        'protocolo': primeira['protocolo'],
        'grouped_items': primeira['grouped_items'],
        'linhas': primeira['linhas'],
        'resumo_publico': primeira['resumo_publico'],
        'total_geral': total_geral,
        'data_emissao_relatorio': timezone.now().date(),
        'layout': layout,
        'usa_layout_publico': convenio_usa_layout_publico(lote.convenio),
        'coluna_terceira_guia': convenio_relatorio_coluna_guia(lote.convenio) if layout == 'publico' else False,
    }


def _cabecalhos_tabela_publico(coluna_guia: bool) -> list[str]:
    cols = ['Data', 'Paciente', 'Associado']
    if coluna_guia:
        cols.append('Número da Guia')
    cols.extend(['Procedimento', 'Modalidade', 'Contraste', 'Valor'])
    return cols


def _linha_tabela_publico(linha: dict, coluna_guia: bool) -> list:
    data = linha.get('data')
    data_fmt = data.strftime('%d/%m/%Y') if data else '—'
    valores = [
        data_fmt,
        linha.get('paciente') or '-',
        linha.get('nome_associado') or '-',
    ]
    if coluna_guia:
        valores.append(linha.get('numero_guia') or '-')
    valor = linha.get('valor') or Decimal('0')
    valores.extend([
        linha.get('procedimento') or '-',
        linha.get('modalidade') or '-',
        'Sim' if linha.get('com_contraste') else 'Não',
        _fmt_moeda_br(valor),
    ])
    return valores


def _fmt_moeda_br(valor):
    from django.utils.numberformat import format as number_format

    try:
        return f'R$ {number_format(valor or 0, decimal_pos=2, force_grouping=True, use_l10n=True)}'
    except (TypeError, ValueError):
        return 'R$ 0,00'


def _excel_borda_azul():
    from openpyxl.styles import Border, Side

    lado = Side(style='thin', color='8EAADB')
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def _excel_escrever_resumo_assinatura(ws, row_start, secao, context, *, incluir_assinatura=False):
    """Bloco RESUMO — LOTE + campos de assinatura (mesmo layout da impressão)."""
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

    resumo = secao.get('resumo_publico')
    if not resumo:
        return row_start

    borda = _excel_borda_azul()
    header_fill = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
    total_qtd_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
    total_valor_fill = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid')
    bold = Font(bold=True)
    center = Alignment(horizontal='center', vertical='center')
    left = Alignment(horizontal='left', vertical='center', wrap_text=True)
    linha_assinatura = Border(bottom=Side(style='thin', color='000000'))

    row = row_start
    lote_id = secao['lote'].id
    assinatura_row = row

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    titulo = ws.cell(row=row, column=1, value=f'RESUMO — LOTE {lote_id}')
    titulo.font = bold
    titulo.fill = header_fill
    titulo.alignment = center
    for col in (1, 2):
        cell = ws.cell(row=row, column=col)
        cell.border = borda
        cell.fill = header_fill
    row += 1

    for item in resumo.get('modalidades') or []:
        c_label = ws.cell(row=row, column=1, value=item.get('label'))
        c_label.font = bold
        c_label.alignment = left
        c_qtd = ws.cell(row=row, column=2, value=item.get('quantidade'))
        c_qtd.alignment = center
        for col in (1, 2):
            ws.cell(row=row, column=col).border = borda
        row += 1

    for label, valor, fill in (
        ('Quantidade total', resumo.get('quantidade_total'), total_qtd_fill),
        ('Valor total', _fmt_moeda_br(resumo.get('valor_total')), total_valor_fill),
    ):
        c_label = ws.cell(row=row, column=1, value=label)
        c_label.font = bold
        c_label.fill = fill
        c_label.alignment = left
        c_valor = ws.cell(row=row, column=2, value=valor)
        c_valor.font = bold
        c_valor.fill = fill
        c_valor.alignment = center
        for col in (1, 2):
            ws.cell(row=row, column=col).border = borda
        row += 1

    resumo_fim = row - 1

    if incluir_assinatura:
        col_rotulo = 4
        col_valor_ini = 5
        col_valor_fim = 7
        empresa = context.get('empresa')
        emissao = context.get('data_emissao_relatorio')
        local = (context.get('local_relatorio') or '').strip()

        campos = [
            ('DATA', emissao.strftime('%d/%m/%Y') if emissao else '', 1),
            ('LOCAL', local, 1),
            ('RESPONSÁVEL PARA ASSINATURA', '', 4),
        ]
        r = assinatura_row
        for rotulo, valor, linhas in campos:
            ws.cell(row=r, column=col_rotulo, value=rotulo).font = bold
            ws.merge_cells(
                start_row=r,
                start_column=col_valor_ini,
                end_row=r + linhas - 1,
                end_column=col_valor_fim,
            )
            val_cell = ws.cell(row=r, column=col_valor_ini, value=valor or None)
            val_cell.alignment = Alignment(vertical='bottom')
            borda_fim = r + linhas - 1
            for col in range(col_valor_ini, col_valor_fim + 1):
                ws.cell(row=borda_fim, column=col).border = linha_assinatura
            r += linhas + 2

        if empresa:
            footer_row = resumo_fim + 2
            cnpj = (empresa.cnpj or '').strip()
            texto = f'{empresa.razao} · CNPJ {cnpj}' if cnpj else empresa.razao
            ws.cell(row=footer_row, column=1, value=texto).font = Font(size=9, color='333333')

        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 14
        ws.column_dimensions['G'].width = 14

    ws.column_dimensions['A'].width = max(ws.column_dimensions['A'].width or 0, 40)
    ws.column_dimensions['B'].width = max(ws.column_dimensions['B'].width or 0, 12)

    return row


def montar_workbook_lote_publico(context):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    wb.remove(wb.active)
    coluna_guia = bool(context.get('coluna_terceira_guia'))
    cabecalhos = _cabecalhos_tabela_publico(coluna_guia)
    header_fill = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
    header_font = Font(bold=True)
    borda = _excel_borda_azul()
    empresa = context.get('empresa')
    convenio = context.get('convenio_nome') or 'Convênio'
    secoes = context.get('secoes') or []
    ultima_secao = len(secoes) - 1

    for idx, secao in enumerate(secoes):
        titulo = f'Lote {secao["lote"].id}' if len(secoes) > 1 else 'Controle Exames'
        ws = wb.create_sheet(title=titulo[:31])
        row = 1
        if empresa:
            ws.cell(row=row, column=1, value=empresa.razao).font = Font(bold=True, size=12)
            row += 1
        ws.cell(row=row, column=1, value=f'CONTROLE DE EXAMES — {convenio}').font = Font(bold=True, size=11)
        row += 1
        ws.cell(row=row, column=1, value=f'MÊS DE REFERÊNCIA: {secao.get("mes_referencia") or "—"}').font = Font(bold=True)
        row += 1
        meta = [
            f'Lote: {secao["lote"].id}',
            f'Protocolo: {secao.get("protocolo") or "—"}',
        ]
        if secao.get('periodo_inicio') and secao.get('periodo_fim'):
            pi = secao['periodo_inicio'].strftime('%d/%m/%Y')
            pf = secao['periodo_fim'].strftime('%d/%m/%Y')
            meta.append(f'Período: {pi} a {pf}')
        emissao = context.get('data_emissao_relatorio')
        if emissao:
            meta.append(f'Emissão: {emissao.strftime("%d/%m/%Y")}')
        ws.cell(row=row, column=1, value=' · '.join(meta))
        row += 2

        for col, titulo_col in enumerate(cabecalhos, start=1):
            cell = ws.cell(row=row, column=col, value=titulo_col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', wrap_text=True)
            cell.border = borda
        row += 1

        for linha in secao.get('linhas') or []:
            for col, valor in enumerate(_linha_tabela_publico(linha, coluna_guia), start=1):
                cell = ws.cell(row=row, column=col, value=valor)
                cell.border = borda
                if col == len(cabecalhos):
                    cell.alignment = Alignment(horizontal='right')
            row += 1

        if secao.get('resumo_publico'):
            row += 1
            row = _excel_escrever_resumo_assinatura(
                ws,
                row,
                secao,
                context,
                incluir_assinatura=(idx == ultima_secao),
            )

        if len(cabecalhos) >= 4:
            ws.column_dimensions['A'].width = max(ws.column_dimensions['A'].width or 0, 12)
            ws.column_dimensions['B'].width = max(ws.column_dimensions['B'].width or 0, 36)
            if coluna_guia:
                ws.column_dimensions['C'].width = 36
                ws.column_dimensions['D'].width = 16
                ws.column_dimensions['E'].width = 42
            else:
                ws.column_dimensions['C'].width = max(ws.column_dimensions['C'].width or 0, 36)
                ws.column_dimensions['D'].width = 42

    if not wb.sheetnames:
        wb.create_sheet('Controle Exames')
    return wb
