"""Importação de produtos de estoque a partir de planilha Excel (.xlsx)."""
from __future__ import annotations

from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook

from estoque.models import ProdutoEstoque

COLUNAS_MODELO = (
    'Código do produto',
    'Descrição do produto',
    'Marca',
    'Quantidade em estoque',
    'Estoque contábil',
    'Estoque mínimo',
    'Dias para entrega',
    'Valor da última compra',
    'Valor custo médio',
    'Valor custo contábil',
)


def _norm_header(val) -> str:
    if val is None:
        return ''
    s = str(val).strip().lower()
    for old, new in (
        ('á', 'a'), ('à', 'a'), ('ã', 'a'), ('â', 'a'),
        ('é', 'e'), ('ê', 'e'), ('í', 'i'),
        ('ó', 'o'), ('ô', 'o'), ('õ', 'o'), ('ú', 'u'), ('ç', 'c'),
    ):
        s = s.replace(old, new)
    s = ' '.join(s.split())
    return s


def _parse_int_opcional(val) -> int | None:
    if val is None or val == '':
        return None
    if isinstance(val, (int, float)):
        return int(val)
    s = str(val).strip()
    if not s:
        return None
    try:
        return int(float(s.replace(',', '.')))
    except (ValueError, TypeError):
        return None


def _parse_decimal(val, *, casas: int = 4) -> Decimal:
    if val is None or val == '':
        return Decimal('0')
    if isinstance(val, (int, float, Decimal)):
        return Decimal(str(val))
    s = str(val).strip().replace('R$', '').replace(' ', '')
    if not s:
        return Decimal('0')
    neg = s.startswith('-')
    s = s.lstrip('-')
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        n = Decimal(s)
    except InvalidOperation:
        return Decimal('0')
    return -n if neg else n


def _map_colunas(headers: list) -> dict[str, int | None]:
    norm = [_norm_header(h) for h in headers]
    aliases = {
        'codigo_produto': (
            'codigo do produto',
            'codigo produto',
            'codigo',
            'cod',
        ),
        'descricao': (
            'descricao do produto',
            'descricao',
            'produto',
        ),
        'marca': ('marca',),
        'quantidade_estoque': (
            'quantidade em estoque',
            'quantidade estoque',
            'qtd estoque',
            'quantidade',
            'qtd',
        ),
        'quantidade_estoque_contabil': (
            'estoque contabil',
            'quantidade estoque contabil',
            'qtd estoque contabil',
        ),
        'valor_ultima_compra': (
            'valor da ultima compra',
            'ultima compra',
            'valor ultima compra',
        ),
        'valor_custo_medio': (
            'valor custo medio',
            'custo medio',
        ),
        'valor_custo_contabil': (
            'valor custo contabil',
            'custo contabil',
        ),
        'estoque_minimo': (
            'estoque minimo',
            'minimo',
            'est minimo',
        ),
        'dias_entrega_fornecedor': (
            'dias para entrega',
            'dias entrega',
            'prazo entrega',
            'lead time',
        ),
    }
    out: dict[str, int | None] = {k: None for k in aliases}
    for campo, opts in aliases.items():
        for i, h in enumerate(norm):
            if h in opts:
                out[campo] = i
                break
    return out


def _celula(row: tuple, idx: int | None):
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _linha_vazia(row: tuple) -> bool:
    return all(v is None or str(v).strip() == '' for v in row)


def parse_planilha_produtos(content: bytes) -> tuple[list[dict], list[str]]:
    wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], ['Planilha vazia.']

    header_idx = None
    mapa: dict[str, int | None] = {}
    for i, row in enumerate(rows[:20]):
        if _linha_vazia(row):
            continue
        candidato = _map_colunas(list(row))
        if candidato['codigo_produto'] is not None and candidato['descricao'] is not None:
            header_idx = i
            mapa = candidato
            break

    if header_idx is None:
        return [], [
            'Cabeçalho não encontrado. Use as colunas do modelo: '
            + ', '.join(COLUNAS_MODELO),
        ]

    erros: list[str] = []
    linhas: list[dict] = []
    for num, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if _linha_vazia(row):
            continue
        codigo = _celula(row, mapa['codigo_produto'])
        descricao = _celula(row, mapa['descricao'])
        codigo_s = str(codigo).strip() if codigo is not None else ''
        descricao_s = str(descricao).strip() if descricao is not None else ''
        if not codigo_s:
            erros.append(f'Linha {num}: código do produto vazio — linha ignorada.')
            continue
        if not descricao_s:
            erros.append(f'Linha {num}: descrição vazia — linha ignorada.')
            continue
        marca = _celula(row, mapa['marca'])
        linhas.append(
            {
                'linha': num,
                'codigo_produto': codigo_s[:50],
                'descricao': descricao_s[:300],
                'marca': (str(marca).strip() if marca is not None else '')[:120],
                'quantidade_estoque': _parse_decimal(_celula(row, mapa['quantidade_estoque'])),
                'quantidade_estoque_contabil': _parse_decimal(
                    _celula(row, mapa['quantidade_estoque_contabil']),
                ),
                'valor_ultima_compra': _parse_decimal(_celula(row, mapa['valor_ultima_compra'])),
                'valor_custo_medio': _parse_decimal(_celula(row, mapa['valor_custo_medio'])),
                'valor_custo_contabil': _parse_decimal(_celula(row, mapa['valor_custo_contabil'])),
                'estoque_minimo': _parse_decimal(_celula(row, mapa['estoque_minimo']))
                if mapa['estoque_minimo'] is not None
                else None,
                'dias_entrega_fornecedor': _parse_int_opcional(
                    _celula(row, mapa['dias_entrega_fornecedor']),
                )
                if mapa['dias_entrega_fornecedor'] is not None
                else None,
            },
        )
    return linhas, erros


@transaction.atomic
def importar_produtos_estoque(
    empresa_id: int,
    linhas: list[dict],
    *,
    atualizar_existentes: bool = True,
) -> dict[str, int]:
    stats = {'criados': 0, 'atualizados': 0, 'ignorados': 0}
    codigos = [ln['codigo_produto'] for ln in linhas]
    existentes = {
        p.codigo_produto: p
        for p in ProdutoEstoque.objects.filter(
            empresa_id=empresa_id,
            codigo_produto__in=codigos,
        )
    }
    criar: list[ProdutoEstoque] = []
    atualizar: list[ProdutoEstoque] = []

    for ln in linhas:
        cod = ln['codigo_produto']
        if cod in existentes:
            if not atualizar_existentes:
                stats['ignorados'] += 1
                continue
            obj = existentes[cod]
            obj.descricao = ln['descricao']
            obj.marca = ln['marca']
            obj.quantidade_estoque = ln['quantidade_estoque']
            obj.quantidade_estoque_contabil = ln['quantidade_estoque_contabil']
            obj.valor_ultima_compra = ln['valor_ultima_compra']
            obj.valor_custo_medio = ln['valor_custo_medio']
            obj.valor_custo_contabil = ln['valor_custo_contabil']
            if ln.get('estoque_minimo') is not None:
                obj.estoque_minimo = ln['estoque_minimo']
            if ln.get('dias_entrega_fornecedor') is not None:
                obj.dias_entrega_fornecedor = ln['dias_entrega_fornecedor']
            atualizar.append(obj)
            stats['atualizados'] += 1
        else:
            criar.append(
                ProdutoEstoque(
                    empresa_id=empresa_id,
                    codigo_produto=cod,
                    descricao=ln['descricao'],
                    marca=ln['marca'],
                    quantidade_estoque=ln['quantidade_estoque'],
                    quantidade_estoque_contabil=ln['quantidade_estoque_contabil'],
                    estoque_minimo=ln['estoque_minimo'] or 0,
                    dias_entrega_fornecedor=ln.get('dias_entrega_fornecedor'),
                    valor_ultima_compra=ln['valor_ultima_compra'],
                    valor_custo_medio=ln['valor_custo_medio'],
                    valor_custo_contabil=ln['valor_custo_contabil'],
                ),
            )
            stats['criados'] += 1

    if criar:
        ProdutoEstoque.objects.bulk_create(criar)
    if atualizar:
        ProdutoEstoque.objects.bulk_update(
            atualizar,
            [
                'descricao',
                'marca',
                'quantidade_estoque',
                'quantidade_estoque_contabil',
                'valor_ultima_compra',
                'valor_custo_medio',
                'valor_custo_contabil',
                'estoque_minimo',
                'dias_entrega_fornecedor',
            ],
        )
    return stats


def resposta_modelo_excel() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Estoque'
    ws.append(list(COLUNAS_MODELO))
    ws.append(['P001', 'Produto exemplo', 'Marca X', 10, 10, 5, 7, 12.5, 11, 11.2])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
