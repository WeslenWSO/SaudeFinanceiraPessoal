from io import BytesIO

from django.contrib.auth.models import User
from django.test import Client, TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from empresa.models import Empresa
from estoque.export_planilha import exportar_produtos_cadastro_xlsx
from estoque.models import ProdutoEstoque


class ExportPlanilhaEstoqueTest(TestCase):
    def test_bytes_contem_codigo(self):
        empresa = Empresa.objects.create(razao='Exp', cnpj='12345678000188')
        p = ProdutoEstoque.objects.create(
            empresa=empresa,
            codigo_produto='X1',
            descricao='Teste',
        )
        data = exportar_produtos_cadastro_xlsx([p])
        wb = load_workbook(BytesIO(data))
        row = list(wb.active.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        self.assertEqual(row[0], 'X1')


@override_settings(ALLOWED_HOSTS=['testserver', 'localhost'])
class ExportExcelViewTest(TestCase):
    def setUp(self):
        self.empresa = Empresa.objects.create(razao='Loja', cnpj='12345678000199')
        self.user = User.objects.create_user('u1', password='secret')
        self.client = Client()
        self.client.login(username='u1', password='secret')
        session = self.client.session
        session['empresa_id'] = self.empresa.pk
        session.save()

    def test_produto_exportar_excel(self):
        ProdutoEstoque.objects.create(
            empresa=self.empresa,
            codigo_produto='WEB',
            descricao='Via HTTP',
        )
        response = self.client.get(reverse('estoque:produto_exportar_excel'))
        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
