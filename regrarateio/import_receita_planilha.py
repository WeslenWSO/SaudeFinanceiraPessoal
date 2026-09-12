"""Importação de receitas (exames) via planilha Excel → contas a receber + rateio."""
from __future__ import annotations

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.db import transaction
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook

from cobranca.models import Cobranca
from contasareceber.models import ContaAReceber
from notasfiscais.models import NotaFiscalServico
from regrarateio.models import LancamentoRateio, RegraRateio, RegraRateioItem
from regrarateio.services import (
    _gerar_linhas_rateio_conta_receber,
    _regra_usa_valor_manual,
    preview_linhas_rateio_por_regra,
)


def _norm_header(val) -> str:
    if val is None:
        return ''
    s = str(val).strip().lower()
    for old, new in (
        ('á', 'a'), ('à', 'a'), ('ã', 'a'), ('â', 'a'),
        ('é', 'e'), ('ê', 'e'), ('í', 'i'),
        ('ó', 'o'), ('ô', 'o'), ('õ', 'o'), ('ú', 'u'), ('ç', 'c'),
    ):
        s = s.replace(old, new)
    return s


def _find_col(norm_headers: list[str], *patterns: str) -> int | None:
    for i, h in enumerate(norm_headers):
        for p in patterns:
            pn = _norm_header(p)
            if h == pn or h.startswith(pn) or pn in h:
                return i
    return None


def _parse_data(val) -> date | None:
    if val is None or val == '':
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    s = str(val).strip()
    if not s:
        return None
    for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            continue
    return None


def _parse_moeda(val) -> Decimal:
    if val is None or val == '':
        return Decimal('0')
    if isinstance(val, (int, float, Decimal)):
        return Decimal(str(val)).quantize(Decimal('0.01'))
    s = str(val).strip().replace('R$', '').replace(' ', '')
    if not s:
        return Decimal('0')
    neg = s.startswith('-')
    s = s.lstrip('-')
    if ',' in s and '.' in s:
        s = s.replace('.', '').replace(',', '.')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        n = Decimal(s)
    except InvalidOperation:
        return Decimal('0')
    return (-n if neg else n).quantize(Decimal('0.01'))


def _map_colunas(headers: list) -> dict[str, int | None]:
    norm = [_norm_header(h) for h in headers]
    return {
        'data': _find_col(norm, 'data'),
        'paciente': _find_col(norm, 'paciente'),
        'procedimento': _find_col(norm, 'procedimento'),
        'modalidade': _find_col(norm, 'modalidade', 'moda'),
        'viabilidade': _find_col(norm, 'viabilidade', 'viabilizacao', 'viabiliz'),
        'forma_pagamento': _find_col(norm, 'forma de pagamento', 'forma pagamento'),
        'valor': _find_col(norm, 'valor'),
        'regra': _find_col(norm, 'regra do rateio', 'regra rateio', 'regra'),
        # compatibilidade modelo antigo
        'socio': _find_col(norm, 'socio'),
        'descricao': _find_col(norm, 'descricao'),
        'nf': _find_col(norm, 'nf', 'nota'),
    }


def _cel(row, col_map, key):
    i = col_map.get(key)
    if i is None or i >= len(row):
        return None
    return row[i]


def _extrair_nf_forma_pagamento(texto: str) -> str:
    m = re.search(r'\bNF\s*(\d+)\b', texto or '', re.I)
    return m.group(1) if m else ''


def _a_faturar(texto: str) -> bool:
    t = (texto or '').strip().upper()
    return 'A FATURAR' in t or t == 'AFATURAR'


def _montar_observacao(paciente: str, procedimento: str, modalidade: str, viabilidade: str) -> str:
    partes = []
    if paciente:
        partes.append(f'Pac:{paciente[:120]}')
    if procedimento:
        partes.append(f'Proc:{procedimento[:160]}')
    if modalidade:
        partes.append(f'Mod:{modalidade}')
    if viabilidade:
        partes.append(f'Viab:{viabilidade}')
    return '|'.join(partes)


def _resolver_regra(empresa_id: int, texto: str, cache: dict) -> RegraRateio | None:
    key = (texto or '').strip().upper()
    if not key:
        return None
    if key in cache:
        return cache[key]
    qs = RegraRateio.objects.filter(empresa_id=empresa_id, rateio='S')
    regra = None
    if key.isdigit():
        regra = qs.filter(pk=int(key)).first()
    if not regra:
        regra = qs.filter(codigo__iexact=texto.strip()).first()
    if not regra:
        regra = qs.filter(nomedaregra__iexact=texto.strip()).first()
    if not regra:
        for r in qs:
            if key in str(r).upper():
                regra = r
                break
    cache[key] = regra
    return regra


def _resolver_cobranca(texto: str, cache: dict) -> Cobranca | None:
    key = (texto or '').strip().upper()
    if not key or _a_faturar(texto):
        return None
    if key in cache:
        return cache[key]
    c = Cobranca.objects.filter(descricao__iexact=texto.strip()).first()
    if not c:
        c = Cobranca.objects.filter(descricao__icontains=texto.strip()[:30]).first()
    cache[key] = c
    return c


def _nota_por_numero(empresa_id: int, nf: str, cache: dict):
    nf = (nf or '').strip()
    if not nf:
        return None
    if nf in cache:
        return cache[nf]
    nota = (
        NotaFiscalServico.objects.filter(empresa_id=empresa_id, numero_nota=nf)
        .order_by('-data_emissao')
        .first()
    )
    cache[nf] = nota
    return nota


def parse_receita_planilha_xlsx(
    file_bytes: bytes,
    *,
    regra_padrao_id: int | None = None,
    empresa_id: int | None = None,
) -> tuple[list[dict], list[str]]:
    wb = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], ['Planilha vazia.']

    col_map = _map_colunas(list(header_row))
    if col_map.get('data') is None:
        return [], ['Coluna «Data» não encontrada.']
    if col_map.get('valor') is None:
        return [], ['Coluna «Valor» não encontrada.']
    if col_map.get('regra') is None and not regra_padrao_id:
        return [], [
            'Informe a coluna «Regra do rateio» na planilha ou selecione uma regra padrão no formulário.'
        ]

    regra_cache: dict = {}
    linhas = []
    avisos = []

    for num, row in enumerate(rows_iter, start=2):
        if not row or not any(row):
            continue

        data = _parse_data(_cel(row, col_map, 'data'))
        valor = _parse_moeda(_cel(row, col_map, 'valor'))
        paciente = str(_cel(row, col_map, 'paciente') or '').strip()
        procedimento = str(_cel(row, col_map, 'procedimento') or '').strip()
        descricao = str(_cel(row, col_map, 'descricao') or '').strip()
        if not procedimento and descricao:
            procedimento = descricao
        if not paciente and not procedimento:
            avisos.append(f'Linha {num}: paciente/procedimento vazio — ignorada.')
            continue
        if not data:
            avisos.append(f'Linha {num}: data inválida — ignorada.')
            continue
        if valor <= 0:
            avisos.append(f'Linha {num}: valor zero ou inválido — ignorada.')
            continue

        modalidade = str(_cel(row, col_map, 'modalidade') or '').strip()
        viabilidade = str(_cel(row, col_map, 'viabilidade') or '').strip()
        forma_pg = str(_cel(row, col_map, 'forma_pagamento') or '').strip()
        nf_col = str(_cel(row, col_map, 'nf') or '').strip()
        regra_txt = str(_cel(row, col_map, 'regra') or '').strip()

        regra_id = regra_padrao_id
        regra_nome = ''
        if regra_txt and empresa_id:
            regra = _resolver_regra(empresa_id, regra_txt, regra_cache)
            if regra:
                regra_id = regra.id
                regra_nome = str(regra)
            else:
                avisos.append(f'Linha {num}: regra «{regra_txt}» não encontrada — ignorada.')
                continue
        elif regra_padrao_id and empresa_id:
            regra = _resolver_regra(empresa_id, str(regra_padrao_id), regra_cache)
            if regra:
                regra_nome = str(regra)

        if not regra_id:
            avisos.append(f'Linha {num}: sem regra de rateio — ignorada.')
            continue

        nf = nf_col or _extrair_nf_forma_pagamento(forma_pg)
        a_faturar = _a_faturar(forma_pg)

        linhas.append(
            {
                'linha': num,
                'data': data.isoformat(),
                'paciente': paciente,
                'procedimento': procedimento,
                'modalidade': modalidade,
                'viabilidade': viabilidade,
                'forma_pagamento': forma_pg,
                'nf': nf,
                'valor': str(valor),
                'regra_id': regra_id,
                'regra_nome': regra_nome,
                'a_faturar': a_faturar,
            }
        )

    if not linhas:
        avisos.append('Nenhuma linha válida encontrada.')
    return linhas, avisos


def validar_linhas_importacao(empresa_id: int, linhas: list[dict]) -> tuple[list[dict], list[str]]:
    erros = []
    regra_cache: dict = {}
    for ln in linhas:
        ln['valido'] = True
        ln['preview_rateio'] = []
        ln['motivo_invalido'] = ''

        if ln.get('a_faturar'):
            ln['valido'] = False
            ln['motivo_invalido'] = 'A FATURAR'
            continue

        regra = _resolver_regra(empresa_id, str(ln['regra_id']), regra_cache)
        if not regra:
            ln['valido'] = False
            ln['motivo_invalido'] = 'Regra inválida'
            erros.append(f"Linha {ln['linha']}: regra de rateio inválida.")
            continue

        itens = list(RegraRateioItem.objects.filter(regrarateio=regra).select_related('socios'))
        if not itens:
            ln['valido'] = False
            ln['motivo_invalido'] = 'Regra sem sócios'
            erros.append(f"Linha {ln['linha']}: regra «{regra}» sem sócios cadastrados.")
            continue
        if _regra_usa_valor_manual(regra, itens):
            ln['valido'] = False
            ln['motivo_invalido'] = 'Regra exige valor manual'
            erros.append(
                f"Linha {ln['linha']}: regra «{regra}» usa valor na aplicação — "
                'use o modal «Contas a receber» para informar valores manuais.'
            )
            continue

        prev = preview_linhas_rateio_por_regra(
            regra.id,
            Decimal(ln['valor']),
            LancamentoRateio.TIPO_RECEBIMENTO,
            empresa_id=empresa_id,
        )
        ln['preview_rateio'] = prev
        ln['regra_nome'] = str(regra)
        if not prev:
            ln['valido'] = False
            ln['motivo_invalido'] = 'Rateio zerado'
            erros.append(f"Linha {ln['linha']}: preview do rateio vazio.")

    return linhas, erros


@transaction.atomic
def importar_receitas_planilha(empresa_id: int, linhas: list[dict]) -> dict:
    regra_cache: dict = {}
    cobranca_cache: dict = {}
    nota_cache: dict = {}

    criados_car = 0
    criados_lr = 0
    ignorados = 0
    erros: list[str] = []

    for ln in linhas:
        if not ln.get('valido', True):
            ignorados += 1
            continue

        data = date.fromisoformat(ln['data'])
        valor = Decimal(ln['valor'])
        paciente = ln.get('paciente') or ''
        procedimento = ln.get('procedimento') or ''
        modalidade = ln.get('modalidade') or ''
        viabilidade = ln.get('viabilidade') or ''
        forma_txt = ln.get('forma_pagamento') or ''
        nf = ln.get('nf') or ''

        regra = _resolver_regra(empresa_id, str(ln['regra_id']), regra_cache)
        if not regra:
            ignorados += 1
            continue

        obs = _montar_observacao(paciente, procedimento, modalidade, viabilidade)
        if ContaAReceber.objects.filter(
            empresa_id=empresa_id,
            data_recebimento=data,
            valor_a_receber=valor,
            observacao=obs,
        ).exists():
            ignorados += 1
            continue

        nota = _nota_por_numero(empresa_id, nf, nota_cache)
        cobranca = _resolver_cobranca(forma_txt, cobranca_cache)
        cliente = (viabilidade or paciente or 'Importação planilha')[:200]

        car = ContaAReceber.objects.create(
            empresa_id=empresa_id,
            nota=nota,
            socio=nota.socio if nota and nota.socio_id else None,
            cliente=cliente,
            data_emissao=data,
            data_vencimento=data,
            data_recebimento=data,
            valor_a_receber=valor,
            valor_recebido=valor,
            status='pago',
            doc=(nf or forma_txt)[:50] or None,
            observacao=obs,
            forma_pagamento=cobranca,
            regra_rateio=regra,
        )
        criados_car += 1

        itens = list(RegraRateioItem.objects.filter(regrarateio=regra).select_related('socios'))
        try:
            n = _gerar_linhas_rateio_conta_receber(car, regra, itens)
            criados_lr += n
        except ValueError as exc:
            car.delete()
            criados_car -= 1
            erros.append(f"Linha {ln['linha']}: {exc}")

    return {
        'criados_car': criados_car,
        'criados_lr': criados_lr,
        'ignorados': ignorados,
        'erros': erros,
    }


def gerar_modelo_receita_planilha_excel() -> HttpResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = 'Receitas'
    ws.append([
        'Data',
        'Paciente',
        'Procedimento',
        'Modalidade',
        'Viabilizacao',
        'Forma de Pagamento',
        'Valor',
        'Regra do rateio',
    ])
    ws.append([
        '03/08/2026',
        'ANTONIO SANTOS GOMES',
        'US - Abdome total',
        'US',
        'Particular',
        'NF 4856 - DH',
        200.00,
        '002 DESPESAS PARTICIPACAO US %',
    ])
    ws.append([
        '03/08/2026',
        'DEBORA ALVES ROGERI TEIXERA',
        'US - Próstata via abdominal',
        'US',
        'FUNCIONAL HEALTH TECH',
        'A FATURAR',
        350.00,
        '002 DESPESAS PARTICIPACAO US %',
    ])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="modelo_receitas_rateio.xlsx"'
    return response
