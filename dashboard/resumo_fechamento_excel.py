"""Exportação do resumo de fechamento para Excel (.xlsx)."""
from __future__ import annotations

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def _parse_moeda_br_txt(valor) -> float:
    if valor is None:
        return 0.0
    s = str(valor).strip().replace(' ', '')
    if not s or s == '—':
        return 0.0
    neg = s.startswith('-')
    s = s.lstrip('-')
    s = s.replace('.', '').replace(',', '.')
    try:
        n = float(s)
    except ValueError:
        return 0.0
    return -n if neg else n


def _estilos():
    header_fill = PatternFill('solid', fgColor='212529')
    header_font = Font(bold=True, color='FFFFFF')
    section_font = Font(bold=True, size=12)
    money_fmt = '#,##0.00'
    thin = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )
    return {
        'header_fill': header_fill,
        'header_font': header_font,
        'section_font': section_font,
        'money_fmt': money_fmt,
        'thin': thin,
        'left': Alignment(horizontal='left', vertical='center', wrap_text=True),
        'right': Alignment(horizontal='right', vertical='center'),
    }


def _cabecalho_ws(ws, ctx, titulo_aba: str, col_count: int, st):
    ws['A1'] = titulo_aba
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = f"Empresa: {ctx.get('empresa_razao_social') or '—'}"
    if ctx.get('empresa_cnpj_fmt'):
        ws['A3'] = f"CNPJ: {ctx['empresa_cnpj_fmt']}"
        ws['A4'] = f"Período: {ctx.get('periodo_titulo', '')}"
        row_meta = 5
    else:
        ws['A3'] = f"Período: {ctx.get('periodo_titulo', '')}"
        row_meta = 4
    if ctx.get('filtro_socio_nome'):
        ws.cell(row=row_meta, column=1, value=f"Sócio(s): {ctx['filtro_socio_nome']}")
        row_meta += 1
    if col_count > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_count)
    return row_meta + 1


def _escrever_tabela(ws, start_row, headers, rows, st, money_cols=None):
    money_cols = money_cols or set()
    thin = st['thin']
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=h)
        cell.fill = st['header_fill']
        cell.font = st['header_font']
        cell.border = thin
        cell.alignment = st['left'] if col == 1 else st['right']
    r = start_row + 1
    for row_data in rows:
        for col, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.border = thin
            if col in money_cols:
                if isinstance(val, (int, float)):
                    cell.number_format = st['money_fmt']
                cell.alignment = st['right']
            else:
                cell.alignment = st['left']
        r += 1
    return r


def _auto_largura(ws, max_col: int, min_width=10, max_width=48):
    for col in range(1, max_col + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, min_width), max_width)


def gerar_resumo_fechamento_excel(ctx) -> HttpResponse:
    st = _estilos()
    wb = Workbook()
    wb.remove(wb.active)

    # --- Aba: Pro-labore ---
    ws_pl = wb.create_sheet('Pro-labore')
    row = _cabecalho_ws(ws_pl, ctx, 'Pró-labore e dividendo', 4, st)
    pl_rows = []
    for linha in ctx.get('grade_prolabore_linhas') or []:
        pl_rows.append([
            linha['socio_nome'],
            _parse_moeda_br_txt(linha['saldo_txt']),
            _parse_moeda_br_txt(linha['pl_value']),
            _parse_moeda_br_txt(linha['dividendo_txt']),
        ])
    if ctx.get('totais_prolabore'):
        tp = ctx['totais_prolabore']
        pl_rows.append([
            'Total',
            _parse_moeda_br_txt(tp['saldo_txt']),
            _parse_moeda_br_txt(tp['prolabore_txt']),
            _parse_moeda_br_txt(tp['dividendo_txt']),
        ])
    row = _escrever_tabela(
        ws_pl,
        row,
        ['Sócio', 'Saldo (consolidado)', 'Pró-labore', 'Dividendo'],
        pl_rows,
        st,
        money_cols={2, 3, 4},
    )
    _auto_largura(ws_pl, 4)

    # --- Aba: Consolidado ---
    ws_cons = wb.create_sheet('Consolidado')
    row = _cabecalho_ws(ws_cons, ctx, 'Resumo consolidado por sócio', 12, st)
    cons_headers = [
        'Sócio', 'Fat. bruto', 'Fat. líquido (NF)', 'Val. recebido', 'Val. a receber',
        'Rec. compet. ant.', 'Total recebido', 'Impostos', 'Desp. rateio',
        'Rec. caixa', 'Pago sócio (dist.)', 'Saldo',
    ]
    cons_rows = []
    for linha in ctx.get('resumo_consolidado_linhas') or []:
        cons_rows.append([
            linha['socio_nome'],
            _parse_moeda_br_txt(linha['b_txt']),
            _parse_moeda_br_txt(linha['c_txt']),
            _parse_moeda_br_txt(linha['d_txt']),
            _parse_moeda_br_txt(linha['e_txt']),
            _parse_moeda_br_txt(linha['f_txt']),
            _parse_moeda_br_txt(linha['g_txt']),
            _parse_moeda_br_txt(linha['h_txt']),
            _parse_moeda_br_txt(linha['i_txt']),
            _parse_moeda_br_txt(linha['j_txt']),
            _parse_moeda_br_txt(linha['l_txt']),
            _parse_moeda_br_txt(linha['k_txt']),
        ])
    if ctx.get('totais_consolidado'):
        tc = ctx['totais_consolidado']
        cons_rows.append([
            'Total',
            _parse_moeda_br_txt(tc.get('b')),
            _parse_moeda_br_txt(tc.get('c')),
            _parse_moeda_br_txt(tc.get('d')),
            _parse_moeda_br_txt(tc.get('e')),
            _parse_moeda_br_txt(tc.get('f')),
            _parse_moeda_br_txt(tc.get('g')),
            _parse_moeda_br_txt(tc.get('h')),
            _parse_moeda_br_txt(tc.get('i')),
            _parse_moeda_br_txt(tc.get('j')),
            _parse_moeda_br_txt(tc.get('l')),
            _parse_moeda_br_txt(tc.get('k')),
        ])
    _escrever_tabela(ws_cons, row, cons_headers, cons_rows, st, money_cols=set(range(2, 13)))
    _auto_largura(ws_cons, 12)

    # --- Aba: Faturamento ---
    ws_fat = wb.create_sheet('Faturamento')
    row = _cabecalho_ws(ws_fat, ctx, 'Faturamento (NFSe)', 12, st)
    fat_headers = [
        'Sócio', 'Fat. bruto', 'Fat. líquido (NF)', 'Soma impostos', 'PIS', 'COFINS',
        'ISS', 'CSLL', 'IRPJ', 'IRPJ adic.', 'Simples', 'Fat. líquido (após imp.)',
    ]
    fat_rows = []
    for linha in ctx.get('grade_linhas') or []:
        fat_rows.append([
            linha['socio_nome'],
            _parse_moeda_br_txt(linha['fat_bruto_txt']),
            _parse_moeda_br_txt(linha['fat_liquido_nf_txt']),
            _parse_moeda_br_txt(linha['imp_soma_txt']),
            _parse_moeda_br_txt(linha['pis_txt']),
            _parse_moeda_br_txt(linha['cofins_txt']),
            _parse_moeda_br_txt(linha['iss_txt']),
            _parse_moeda_br_txt(linha['csll_txt']),
            _parse_moeda_br_txt(linha['irpj_txt']),
            _parse_moeda_br_txt(linha['irpj_ad_txt']),
            _parse_moeda_br_txt(linha['simples_txt']),
            _parse_moeda_br_txt(linha['fat_liquido_final_txt']),
        ])
    if ctx.get('totais_linha'):
        tl = ctx['totais_linha']
        fat_rows.append([
            'Total',
            _parse_moeda_br_txt(tl['fat_bruto_txt']),
            _parse_moeda_br_txt(tl['fat_liquido_nf_txt']),
            _parse_moeda_br_txt(tl['imp_soma_txt']),
            _parse_moeda_br_txt(tl['pis_txt']),
            _parse_moeda_br_txt(tl['cofins_txt']),
            _parse_moeda_br_txt(tl['iss_txt']),
            _parse_moeda_br_txt(tl['csll_txt']),
            _parse_moeda_br_txt(tl['irpj_txt']),
            _parse_moeda_br_txt(tl['irpj_ad_txt']),
            _parse_moeda_br_txt(tl['simples_txt']),
            _parse_moeda_br_txt(tl['fat_liquido_final_txt']),
        ])
    _escrever_tabela(ws_fat, row, fat_headers, fat_rows, st, money_cols=set(range(2, 13)))
    _auto_largura(ws_fat, 12)

    # --- Aba: Rateio ---
    ws_rat = wb.create_sheet('Rateio')
    row = _cabecalho_ws(ws_rat, ctx, 'Rateio — pagamentos', 7, st)
    rat_rows = []
    for linha in ctx.get('grade_rateio_pagamento') or []:
        rat_rows.append([
            linha['data_txt'],
            linha['emissao_txt'],
            linha['nota_txt'],
            _parse_moeda_br_txt(linha['titulo_valor_txt']),
            linha['descricao'],
            linha['socio_nome'],
            _parse_moeda_br_txt(linha['valor_txt']),
        ])
    if ctx.get('rateio_pg_total_txt'):
        rat_rows.append(['', '', '', '', '', 'Total', _parse_moeda_br_txt(ctx['rateio_pg_total_txt'])])
    _escrever_tabela(
        ws_rat,
        row,
        ['Data pagamento', 'Data emissão', 'Nota', 'Valor título', 'Descrição', 'Sócio', 'Valor rateio'],
        rat_rows,
        st,
        money_cols={4, 7},
    )
    _auto_largura(ws_rat, 7)

    # --- Aba: Contas recebidas ---
    ws_cr = wb.create_sheet('Contas recebidas')
    row = _cabecalho_ws(ws_cr, ctx, 'Contas recebidas no período', 9, st)
    cr_rows = []
    for linha in ctx.get('contas_recebidas_periodo') or []:
        cr_rows.append([
            linha['id'],
            linha['cliente'],
            linha['emissao_txt'],
            linha['recebimento_txt'],
            linha['nota_txt'],
            linha['socio_txt'],
            linha['parcela'],
            linha['doc'],
            _parse_moeda_br_txt(linha['valor_txt']),
        ])
    if ctx.get('contas_recebidas_total_txt'):
        cr_rows.append(['', '', '', '', '', '', '', 'Total', _parse_moeda_br_txt(ctx['contas_recebidas_total_txt'])])
    _escrever_tabela(
        ws_cr,
        row,
        ['ID', 'Cliente', 'Emissão', 'Recebimento', 'NF', 'Sócio', 'Parc.', 'Doc.', 'Valor recebido'],
        cr_rows,
        st,
        money_cols={9},
    )
    _auto_largura(ws_cr, 9)

    # --- Aba: Sem rateio ---
    ws_sr = wb.create_sheet('Sem rateio')
    row = _cabecalho_ws(ws_sr, ctx, 'Pagamentos sem rateio gerado', 8, st)
    sr_rows = []
    for linha in ctx.get('pagamentos_sem_rateio_pagar') or []:
        sr_rows.append([
            linha['id_titulo'],
            linha['emissao_txt'],
            linha['data_txt'],
            linha['contraparte'],
            linha['descricao'],
            _parse_moeda_br_txt(linha['valor_txt']),
            linha['regra'],
            linha['motivo'],
        ])
    _escrever_tabela(
        ws_sr,
        row,
        ['ID', 'Emissão', 'Pagamento', 'Fornecedor', 'Descrição', 'Valor base', 'Regra', 'Situação'],
        sr_rows,
        st,
        money_cols={6},
    )
    _auto_largura(ws_sr, 8)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    di = (ctx.get('data_inicio') or '').replace('-', '')
    df = (ctx.get('data_fim') or '').replace('-', '')
    filename = f'resumo_fechamento_{di}_{df}.xlsx'

    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
