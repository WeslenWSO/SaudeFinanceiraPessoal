"""Exportação do resumo fechamento por resultado para Excel (.xlsx)."""
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font

from dashboard.resumo_fechamento_excel import (
    _auto_largura,
    _cabecalho_ws,
    _escrever_tabela,
    _estilos,
    _parse_moeda_br_txt,
)


def _linhas_grade(linhas, *, com_receita_extra=False):
    rows = []
    for linha in linhas or []:
        base = [
            linha['data_txt'],
            linha['emissao_txt'],
            linha['nota_txt'],
            linha['descricao'],
            linha['socio_nome'],
        ]
        if com_receita_extra:
            base.extend([
                linha.get('modalidade_txt') or '—',
                linha.get('forma_pagamento_txt') or '—',
                linha.get('viabilidade_txt') or '—',
            ])
        base.extend([
            _parse_moeda_br_txt(linha['titulo_valor_txt']),
            _parse_moeda_br_txt(linha['valor_txt']),
        ])
        rows.append(base)
    return rows


def _decimal_mes(val) -> Decimal:
    if isinstance(val, Decimal):
        return val
    if val is None:
        return Decimal('0')
    return Decimal(str(val))


def _linhas_resumo_mensal(ctx) -> list[list]:
    """Monta linhas mês a mês + distribuição por sócio (% informados na tela)."""
    meses = ctx.get('resumo_mensal') or []
    dist = ctx.get('distribuicao_resultado') or []
    rows: list[list] = []
    for mes in meses:
        receita = _decimal_mes(mes.get('receita'))
        despesa = _decimal_mes(mes.get('despesa'))
        resultado = _decimal_mes(mes.get('resultado'))
        rows.append([
            mes.get('label') or '—',
            float(receita),
            float(despesa),
            float(resultado),
            '',
            '',
        ])
        for item in dist:
            pct = Decimal(str(item.get('pct') or 0))
            nome = (item.get('nome') or '').strip()
            if not nome or pct <= 0:
                continue
            valor_soc = (resultado * pct / Decimal('100')).quantize(Decimal('0.01'))
            rows.append(['', '', '', '', nome, float(valor_soc)])
    if not rows and meses:
        rows.append(['—', 0, 0, 0, '(configure % na tela)', 0])
    return rows


def _linhas_distribuicao_periodo(ctx) -> list[list]:
    dist = ctx.get('distribuicao_resultado') or []
    resultado = _decimal_mes(ctx.get('resultado_decimal'))
    rows = []
    for item in dist:
        pct = _decimal_mes(item.get('pct'))
        nome = (item.get('nome') or '').strip()
        if not nome or pct <= 0:
            continue
        val = (resultado * pct / Decimal('100')).quantize(Decimal('0.01'))
        rows.append([nome, float(pct), float(val)])
    return rows


def gerar_resumo_fechamento_resultado_excel(ctx) -> HttpResponse:
    st = _estilos()
    wb = Workbook()
    wb.remove(wb.active)

    # --- Aba: Resumo ---
    ws_res = wb.create_sheet('Resumo')
    row = _cabecalho_ws(ws_res, ctx, 'Resumo do resultado', 6, st)
    resumo_rows = [
        ['Total receitas rateadas', _parse_moeda_br_txt(ctx.get('total_receita_txt'))],
        ['Total despesas rateadas', _parse_moeda_br_txt(ctx.get('total_despesa_txt'))],
        ['Resultado', _parse_moeda_br_txt(ctx.get('resultado_txt'))],
    ]
    row = _escrever_tabela(ws_res, row, ['Item', 'Valor (R$)'], resumo_rows, st, money_cols={2})
    cell_res = ws_res.cell(row=row - 1, column=1)
    cell_res.font = Font(bold=True)
    ws_res.cell(row=row - 1, column=2).font = Font(bold=True)

    dist_rows = _linhas_distribuicao_periodo(ctx)
    if dist_rows:
        row += 1
        ws_res.cell(row=row, column=1, value='Distribuição do resultado (período)').font = st['section_font']
        row += 1
        row = _escrever_tabela(
            ws_res,
            row,
            ['Sócio', '% do resultado', 'Valor (R$)'],
            dist_rows,
            st,
            money_cols={3},
        )

    row += 1
    ws_res.cell(row=row, column=1, value='Resultado por mês').font = st['section_font']
    row += 1
    mes_headers = ['Mês', 'Receita (R$)', 'Despesas (R$)', 'Resultado (R$)', 'Sócio', 'Valor (R$)']
    mes_rows = _linhas_resumo_mensal(ctx)
    row = _escrever_tabela(ws_res, row, mes_headers, mes_rows, st, money_cols={2, 3, 4, 6})
    _auto_largura(ws_res, 6)

    grade_headers_desp = [
        'Data',
        'Emissão',
        'NF',
        'Descrição',
        'Sócio',
        'Valor título',
        'Valor rateio',
    ]
    grade_headers_rec = [
        'Data',
        'Emissão',
        'NF',
        'Descrição',
        'Sócio',
        'Modalidade',
        'Forma de pagamento',
        'Viabilidade',
        'Valor título',
        'Valor rateio',
    ]

    # --- Aba: Receitas ---
    ws_rec = wb.create_sheet('Receitas rateadas')
    row = _cabecalho_ws(ws_rec, ctx, 'Receitas rateadas', 10, st)
    rec_rows = _linhas_grade(ctx.get('linhas_receita'), com_receita_extra=True)
    if ctx.get('total_receita_txt'):
        rec_rows.append(
            ['', '', '', '', '', '', '', '', 'Total', _parse_moeda_br_txt(ctx['total_receita_txt'])]
        )
    _escrever_tabela(ws_rec, row, grade_headers_rec, rec_rows, st, money_cols={9, 10})
    _auto_largura(ws_rec, 10)

    # --- Aba: Despesas ---
    ws_desp = wb.create_sheet('Despesas rateadas')
    row = _cabecalho_ws(ws_desp, ctx, 'Despesas rateadas', 7, st)
    desp_headers = ['Data pgto'] + grade_headers_desp[1:]
    desp_rows = _linhas_grade(ctx.get('linhas_despesa'))
    if ctx.get('total_despesa_txt'):
        desp_rows.append(['', '', '', '', '', 'Total', _parse_moeda_br_txt(ctx['total_despesa_txt'])])
    _escrever_tabela(ws_desp, row, desp_headers, desp_rows, st, money_cols={6, 7})
    _auto_largura(ws_desp, 7)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    di = (ctx.get('data_inicio') or '').replace('-', '')
    df = (ctx.get('data_fim') or '').replace('-', '')
    filename = f'resumo_fechamento_resultado_{di}_{df}.xlsx'

    response = HttpResponse(
        buf.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response
